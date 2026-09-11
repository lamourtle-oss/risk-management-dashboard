"""Turn Excel workbooks into dashboard.json for the 5-tab STL dashboard."""

from __future__ import annotations

import hashlib
import json
import os
import re
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DATA = DOCS / "data"
BANGKOK = timezone(timedelta(hours=7))
OIL_NAMES = {"น้ำมัน", "oil", "fuel"}


def _norm(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").strip().lower())


def _num(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").replace("%", "").strip()
    try:
        return float(text)
    except ValueError:
        return 0.0


def _row_map(headers: tuple[Any, ...], row: tuple[Any, ...]) -> dict[str, Any]:
    data: dict[str, Any] = {}
    for idx, header in enumerate(headers):
        key = _norm(header)
        if not key:
            continue
        data[key] = row[idx] if idx < len(row) else ""
    return data


def _pick(row: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        token = _norm(key)
        if token in row and row[token] not in (None, ""):
            return row[token]
        for actual, value in row.items():
            if token and token in actual and value not in (None, ""):
                return value
    return ""


def read_sheets(path: Path) -> dict[str, list[dict[str, Any]]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    sheets: dict[str, list[dict[str, Any]]] = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = rows[0]
        items = []
        for raw in rows[1:]:
            if not raw or all(v is None or str(v).strip() == "" for v in raw):
                continue
            items.append(_row_map(headers, raw))
        sheets[_norm(ws.title)] = items
    wb.close()
    return sheets


def _find(sheets: dict[str, list[dict[str, Any]]], *needles: str) -> list[dict[str, Any]]:
    wanted = [_norm(n) for n in needles]
    for name, rows in sheets.items():
        if any(w and w in name for w in wanted):
            return rows
    return []


def oil_bucket(liters: float) -> str:
    if liters <= 50:
        return "0-50 ลิตร"
    if liters <= 100:
        return "51-100 ลิตร"
    if liters <= 150:
        return "101-150 ลิตร"
    if liters <= 200:
        return "151-200 ลิตร"
    return "มากกว่า 200 ลิตร"


def build_payload(all_sheets: dict[str, list[dict[str, Any]]], source: str, files: list[str]) -> dict[str, Any]:
    overview = []
    for row in _find(all_sheets, "nplภาพรวม", "nploverview"):
        overview.append(
            {
                "quarter": str(_pick(row, "ไตรมาส", "quarter")),
                "type": str(_pick(row, "ประเภท NPL", "ประเภท", "type")),
                "ratio": _num(_pick(row, "สัดส่วน NPL (%)", "สัดส่วน", "ratio")),
                "amount": _num(_pick(row, "มูลค่า NPL", "มูลค่า", "amount")),
            }
        )

    stages = []
    for row in _find(all_sheets, "nplstage"):
        stages.append(
            {
                "quarter": str(_pick(row, "ไตรมาส")),
                "type": str(_pick(row, "ประเภท NPL", "ประเภท")),
                "stage": str(_pick(row, "Stage", "stage")),
                "contracts": int(_num(_pick(row, "จำนวนสัญญา", "สัญญา"))),
                "amount": _num(_pick(row, "มูลค่า")),
            }
        )

    closures = []
    for row in _find(all_sheets, "nplปิดสัญญา", "ปิดสัญญา"):
        closures.append(
            {
                "quarter": str(_pick(row, "ไตรมาส")),
                "loanType": str(_pick(row, "ประเภทสินเชื่อ", "ประเภท")),
                "writeoffContracts": int(_num(_pick(row, "จำนวน Writeoff"))),
                "writeoffAmount": _num(_pick(row, "มูลค่า Writeoff")),
                "repossessContracts": int(_num(_pick(row, "จำนวนยึดหลักประกัน"))),
                "repossessAmount": _num(_pick(row, "มูลค่ายึดหลักประกัน")),
                "normalContracts": int(_num(_pick(row, "จำนวนปิดปกติ"))),
                "normalAmount": _num(_pick(row, "มูลค่าปิดปกติ")),
            }
        )

    hp_home = []
    for row in _find(all_sheets, "nplhphome", "hphome2568"):
        hp_home.append(
            {
                "month": str(_pick(row, "เดือน", "month")),
                "ratio": _num(_pick(row, "สัดส่วน NPL (%)", "สัดส่วน")),
                "amount": _num(_pick(row, "มูลค่า NPL", "มูลค่า")),
            }
        )

    mix = []
    for row in _find(all_sheets, "ยอดขายสัดส่วน", "salesmix"):
        category = str(_pick(row, "ประเภทสินค้า", "สินค้า", "category"))
        if _norm(category) in OIL_NAMES:
            continue
        mix.append(
            {
                "quarter": str(_pick(row, "ไตรมาส")),
                "category": category,
                "amount": _num(_pick(row, "ยอดขาย", "amount")),
            }
        )

    by_branch = []
    for row in _find(all_sheets, "ยอดขายสาขา", "salesbranch"):
        by_branch.append(
            {
                "quarter": str(_pick(row, "ไตรมาส")),
                "branch": str(_pick(row, "สาขา", "branch")),
                "province": str(_pick(row, "จังหวัด", "province")),
                "amount": _num(_pick(row, "ยอดขาย")),
            }
        )

    by_product = []
    for row in _find(all_sheets, "ยอดขายสินค้า", "salesproduct"):
        product = str(_pick(row, "สินค้า", "product"))
        if _norm(product) in OIL_NAMES:
            continue
        by_product.append(
            {
                "quarter": str(_pick(row, "ไตรมาส")),
                "branch": str(_pick(row, "สาขา")),
                "province": str(_pick(row, "จังหวัด")),
                "product": product,
                "amount": _num(_pick(row, "ยอดขาย")),
            }
        )

    aging = []
    for row in _find(all_sheets, "aging", "อายุสินค้า"):
        aging.append(
            {
                "product": str(_pick(row, "สินค้า")),
                "grade": str(_pick(row, "สภาพ", "grade")),
                "cost": _num(_pick(row, "Cost", "ต้นทุน")),
                "provision": _num(_pick(row, "Provision", "สำรอง")),
                "age0to2": int(_num(_pick(row, "0-2 ปี"))),
                "age3to4": int(_num(_pick(row, "3-4 ปี"))),
                "age4plus": int(_num(_pick(row, "4 ปีขึ้นไป"))),
                "location": str(_pick(row, "Location", "ที่ตั้ง")),
                "branch": str(_pick(row, "สาขา")),
                "province": str(_pick(row, "จังหวัด")),
            }
        )

    minmax = []
    for row in _find(all_sheets, "minmax", "stock"):
        minmax.append(
            {
                "branch": str(_pick(row, "สาขา")),
                "province": str(_pick(row, "จังหวัด")),
                "minStock": _num(_pick(row, "Min stock", "min")),
                "maxStock": _num(_pick(row, "Max stock", "max")),
                "onHand": _num(_pick(row, "คงเหลือ", "onhand")),
            }
        )

    oil = []
    for row in _find(all_sheets, "น้ำมัน", "oil"):
        liters = _num(_pick(row, "ปริมาณลิตร", "ลิตร", "liters"))
        oil.append(
            {
                "quarter": str(_pick(row, "ไตรมาส")),
                "branch": str(_pick(row, "สาขา")),
                "province": str(_pick(row, "จังหวัด")),
                "liters": liters,
                "bucket": oil_bucket(liters),
            }
        )

    sla = []
    for row in _find(all_sheets, "slaซ่อม", "ระยะเวลาซ่อม"):
        sla.append(
            {
                "quarter": str(_pick(row, "ไตรมาส")),
                "bucket": str(_pick(row, "ระยะเวลาซ่อม", "bucket")),
                "jobs": int(_num(_pick(row, "จำนวนงาน", "jobs"))),
            }
        )

    provinces = []
    for row in _find(all_sheets, "slajังหวัด", "slaจังหวัด"):
        provinces.append(
            {
                "quarter": str(_pick(row, "ไตรมาส")),
                "province": str(_pick(row, "จังหวัด")),
                "branch": str(_pick(row, "สาขา")),
                "jobs": int(_num(_pick(row, "งานทั้งหมด"))),
                "slaMet": int(_num(_pick(row, "ตาม SLA"))),
                "slaMiss": int(_num(_pick(row, "ไม่ตาม SLA"))),
            }
        )

    csat = []
    for row in _find(all_sheets, "csat", "ความพึงพอใจ"):
        csat.append(
            {
                "quarter": str(_pick(row, "ไตรมาส")),
                "score": _num(_pick(row, "คะแนนความพึงพอใจ", "คะแนน", "score")),
            }
        )

    quarters = []
    for collection in (overview, mix, by_branch, oil, sla):
        for item in collection:
            q = item.get("quarter")
            if q and q not in quarters:
                quarters.append(q)

    now = datetime.now(BANGKOK)
    source_labels = {
        "sharepoint": "SharePoint · IT Audit / Risk Management",
        "sharepoint-local": "SharePoint · โฟลเดอร์ซิงค์บนเครื่อง",
        "inbox": "inbox/",
        "sample": "ข้อมูลตัวอย่างตามไฟล์สเปก",
    }
    return {
        "generatedAt": now.isoformat(timespec="seconds"),
        "generatedAtLabel": now.strftime("%d/%m/%Y %H:%M น."),
        "source": source,
        "sourceLabel": source_labels.get(source, source),
        "files": files,
        "quarters": quarters,
        "npl": {"overview": overview, "byStage": stages, "closures": closures, "hpHomeNew": hp_home},
        "sales": {"mix": mix, "byBranch": by_branch, "byProduct": by_product},
        "inventory": {"aging": aging, "minmax": minmax},
        "oil": oil,
        "service": {"sla": sla, "provinces": provinces, "csat": csat},
    }


def write_config() -> None:
    password = os.environ.get("DASHBOARD_PASSWORD") or "Risk2026"
    digest = hashlib.sha256(password.encode("utf-8")).hexdigest()
    repo = os.environ.get("GITHUB_REPO", "").strip()
    (DOCS / "js").mkdir(parents=True, exist_ok=True)
    (DOCS / "js" / "config.js").write_text(
        "window.APP_CONFIG = {\n"
        f'  passwordHash: "{digest}",\n'
        f'  githubRepo: "{repo}"\n'
        "};\n",
        encoding="utf-8",
    )


def process(input_dir: Path, source: str) -> Path:
    DATA.mkdir(parents=True, exist_ok=True)
    excel_files = sorted(
        p
        for p in input_dir.rglob("*")
        if p.is_file() and p.suffix.lower() in {".xlsx", ".xls", ".xlsm"} and not p.name.startswith("~$")
    )
    merged: dict[str, list[dict[str, Any]]] = {}
    for path in excel_files:
        for name, rows in read_sheets(path).items():
            merged.setdefault(name, []).extend(rows)
    files = [str(p.relative_to(input_dir)).replace("\\", "/") for p in excel_files]
    payload = build_payload(merged, source, files)
    out = DATA / "dashboard.json"
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    write_config()
    print(f"Wrote {out}")
    return out
